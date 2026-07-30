# -*- coding: utf-8 -*-
"""
Une Formato regla-flex mm.xlsx + Formato regla-flex in.xlsx en Formato Regla Flex.xlsm:

- Base: mm (fórmulas duales más correctas).
- Selector mm/in en J11; alcance F10; división J9.
- Intervalos A/B siguen $F$10; VLOOKUPs duales con IFERROR; sin #N/A.
- Power Query AGD + clientes + patrones (AG-004 / AG-012).
- Certificado D4=AGD / E4=número / F4=año.
- Elimina AG-ADM-F10-00.
- Macros: Guardar, mm/in, Actualizar, Ir a Portada.
"""
from __future__ import annotations

import re
import shutil
import sys
import time
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace", line_buffering=True)

import pythoncom
import win32com.client

FOLDER = Path(r"C:\Users\AG\Desktop\FORMATOS AG")
SOURCE_MM = FOLDER / "Formato regla-flex mm.xlsx"
SOURCE_IN = FOLDER / "Formato regla-flex in.xlsx"
TARGET = FOLDER / "Formato Regla Flex.xlsm"
PASSWORD = "AG-Calidad-2026"
MSO_ROUNDED = 5
XL_XLSM = 52
XL_VERY_HIDDEN = 2

API_BASE = (
    "https://us-central1-agg1-b7f40.cloudfunctions.net/obtenerDatosExcel"
    "?key=TU_CLAVE_SECRETA_AG_APP_2026"
)

HIST_SHEET = "obtenerDatosExcel"
CLIENTES_SHEET = "BD_Clientes"
PATRONES_SHEET = "BD_Patrones"

HIST_COLUMNS = [
    "Name", "certificado", "cliente", "equipo", "marca", "modelo", "serie", "id",
    "fecha", "tecnico", "lugarCalibracion", "frecuenciaCalibracion", "fechaRecepcion",
]
CLIENT_COLUMNS = ["Nombre", "Domicilio", "Contacto", "Correo", "Telefono"]
PATRON_COLUMNS = [
    "noControl", "descripcion", "marca", "modelo", "serie", "noCertificado",
    "fechaUltimaCalibracion", "fechaVencimiento", "estadoProceso", "statusVigencia",
    "laboratorio",
]

DIM_IDS = [
    "AG-001", "AG-002", "AG-004", "AG-006", "AG-007", "AG-012",
    "AG-029", "AG-030", "AG-041", "AG-059",
]


def m_list(values: list[str]) -> str:
    return "{" + ", ".join(f'"{v}"' for v in values) + "}"


M_HISTORIAL = f'''let
    Url = "{API_BASE}",
    Fuente = Json.Document(Web.Contents(Url)),
    HistorialLista = Fuente[historial],
    ClientesLista = Fuente[clientes],
    HistorialBase = Table.FromRecords(HistorialLista, {m_list(HIST_COLUMNS)}, MissingField.UseNull),
    HistorialAGD = Table.SelectRows(
        HistorialBase,
        each Text.StartsWith(Text.Upper(Text.From([certificado])), "AGD-")
    ),
    ClientesBase = Table.FromRecords(ClientesLista, {m_list(CLIENT_COLUMNS)}, MissingField.UseNull),
    NormalizarNombre = (valor as any) as text =>
        let
            Texto = Text.Upper(Text.Trim(if valor = null then "" else Text.From(valor))),
            SinParentesis = if Text.Contains(Texto, "(") then Text.BeforeDelimiter(Texto, "(") else Texto,
            SinAcentos = List.Accumulate(
                {{{{"Á", "A"}}, {{"É", "E"}}, {{"Í", "I"}}, {{"Ó", "O"}}, {{"Ú", "U"}}, {{"Ü", "U"}}, {{"Ñ", "N"}}}},
                SinParentesis,
                (estado, par) => Text.Replace(estado, par{{0}}, par{{1}})
            ),
            Permitidos = Text.Select(SinAcentos, {{"A".."Z", "0".."9", " "}}),
            Compacto = Text.Combine(List.Select(Text.Split(Permitidos, " "), each _ <> ""), " ")
        in
            Compacto,
    HistorialClave = Table.AddColumn(
        HistorialAGD, "_clienteKey", each NormalizarNombre([cliente]), type text
    ),
    ClientesClave = Table.AddColumn(
        ClientesBase, "_clienteKey", each NormalizarNombre([Nombre]), type text
    ),
    ClientesUnicos = Table.Group(
        ClientesClave,
        {{"_clienteKey"}},
        {{{{"ClienteRow", each Table.First(_), type record}}}}
    ),
    Cruzado = Table.NestedJoin(
        HistorialClave, {{"_clienteKey"}},
        ClientesUnicos, {{"_clienteKey"}},
        "ClienteMatch", JoinKind.LeftOuter
    ),
    ClienteRow = Table.AddColumn(
        Cruzado,
        "_clienteRow",
        each if Table.IsEmpty([ClienteMatch]) then null else [ClienteMatch]{{0}}[ClienteRow],
        type nullable record
    ),
    Domicilio = Table.AddColumn(
        ClienteRow, "domicilio",
        each if [_clienteRow] = null then "" else Record.FieldOrDefault([_clienteRow], "Domicilio", ""),
        type text
    ),
    Contacto = Table.AddColumn(
        Domicilio, "contacto",
        each if [_clienteRow] = null then "" else Record.FieldOrDefault([_clienteRow], "Contacto", ""),
        type text
    ),
    Correo = Table.AddColumn(
        Contacto, "correo",
        each if [_clienteRow] = null then "" else Record.FieldOrDefault([_clienteRow], "Correo", ""),
        type text
    ),
    Telefono = Table.AddColumn(
        Correo, "telefono",
        each if [_clienteRow] = null then "" else Record.FieldOrDefault([_clienteRow], "Telefono", ""),
        type text
    ),
    Resultado = Table.SelectColumns(
        Telefono,
        {m_list(HIST_COLUMNS + ["domicilio", "contacto", "correo", "telefono"])}
    )
in
    Resultado'''

M_CLIENTES = f'''let
    Url = "{API_BASE}&formato=clientes",
    Fuente = Json.Document(Web.Contents(Url)),
    Tabla = Table.FromRecords(Fuente, {m_list(CLIENT_COLUMNS)}, MissingField.UseNull),
    Limpio = Table.TransformColumns(
        Tabla,
        List.Transform(
            {m_list(CLIENT_COLUMNS)},
            each {{_, (v) => if v = null then "" else Text.Trim(Text.From(v)), type text}}
        )
    )
in
    Limpio'''

M_PATRONES = f'''let
    Url = "{API_BASE}&formato=patrones",
    Fuente = Json.Document(Web.Contents(Url)),
    Tabla = Table.FromRecords(Fuente, {m_list(PATRON_COLUMNS)}, MissingField.UseNull),
    SoloDim = Table.SelectRows(
        Tabla,
        each let
            d = Text.Upper(if [descripcion] = null then "" else Text.From([descripcion])),
            id = Text.Upper(Text.Trim(if [noControl] = null then "" else Text.From([noControl])))
        in
            Text.Contains(d, "BLOQUE") or Text.Contains(d, "GAUGE")
            or Text.Contains(d, "REGLA") or Text.Contains(d, "RETICULA")
            or Text.Contains(d, "FLEX") or Text.Contains(d, "CINTA")
            or List.Contains({m_list(DIM_IDS)}, id)
    ),
    Fechas = Table.TransformColumns(
        SoloDim,
        {{
            {{"fechaUltimaCalibracion", each try Date.FromText(Text.Start(Text.From(_), 10)) otherwise null, type date}},
            {{"fechaVencimiento", each try Date.FromText(Text.Start(Text.From(_), 10)) otherwise null, type date}}
        }}
    )
in
    Fechas'''

HS = HIST_SHEET
CERT_KEY = 'TRIM($D$4)&"-"&TEXT($E$4,"0000")&"-"&TEXT($F$4,"00")'
MATCH = f"MATCH({CERT_KEY},{HS}!$B:$B,0)"


def idx(col: str, blank: str = '""') -> str:
    return f"=IFERROR(INDEX({HS}!${col}:${col},{MATCH}),{blank})"


# Layout Calculos regla-flex
CALC_FORMULAS = {
    "B5": idx("C"),
    "E5": idx("P"),
    "B6": idx("N"),
    "E6": idx("Q"),
    "B7": idx("O"),
    "E7": '=""',
    "B9": idx("D", '"No encontrado"'),
    "F9": idx("H"),
    "B10": idx("E", '"No encontrado"'),
    "B11": idx("F"),
    "B12": idx("G"),
    "I4": (
        "=IFERROR("
        f'IF(INDEX({HS}!$M:$M,{MATCH})="",'
        f'IF(UPPER(LEFT(INDEX({HS}!$K:$K,{MATCH}),1))="S","Servicio en Sitio",""),'
        f"VALUE(INDEX({HS}!$M:$M,{MATCH}))),"
        f'IF(IFERROR(UPPER(LEFT(INDEX({HS}!$K:$K,{MATCH}),1)),"")="S","Servicio en Sitio",""))'
    ),
    "I5": f'=IFERROR(VALUE(INDEX({HS}!$I:$I,{MATCH})),"")',
    "I6": (
        f'=IFERROR(EDATE($I$5,IF(INDEX({HS}!$L:$L,{MATCH})="6 meses",6,'
        f'IF(INDEX({HS}!$L:$L,{MATCH})="3 meses",3,'
        f'IF(INDEX({HS}!$L:$L,{MATCH})="24 meses",24,'
        f'IF(OR(INDEX({HS}!$L:$L,{MATCH})="12 meses",INDEX({HS}!$L:$L,{MATCH})="1 año"),12,'
        f'IF(ISNUMBER($I$2),$I$2,12)))))),"")'
    ),
    "I7": "=TODAY()",
    "AH4": idx("K"),
    "C14": (
        '=IF(OR($AH$4="Laboratorio",$AH$4="laboratorio"),"Instalaciones AG",'
        'IF(OR($AH$4="Sitio",$AH$4="sitio",'
        'UPPER(LEFT($AH$4,1))="S"),"Instalaciones de Cliente",'
        'IF(OR($I$4="Servicio en Sitio",$I$4="Servicio en sitio"),'
        '"Instalaciones de Cliente","Instalaciones AG")))'
    ),
    "M8": idx("J"),
}

F_CERT = '=Calculos!D4&"-"&TEXT(Calculos!E4,"0000")&"-"&TEXT(Calculos!F4,"00")'

PATRON_BLOCKS = [
    ("Patrones", "D4", "D5", "D7"),    # AG-004 Regla
    ("Patrones", "D18", "D19", "D21"),  # AG-012 Reticula
]


def patron_lookup(id_cell: str, col: str) -> str:
    return (
        f"INDEX({PATRONES_SHEET}!${col}:${col},"
        f"MATCH(TRIM(${id_cell[0]}${id_cell[1:]}),{PATRONES_SHEET}!$A:$A,0))"
    )


def wire_patron_blocks(wb) -> int:
    wired = 0
    for sheet_name, id_cell, cert_cell, vig_cell in PATRON_BLOCKS:
        ws = wb.Worksheets(sheet_name)
        pid = str(ws.Range(id_cell).Value or "").strip()
        if not pid:
            continue
        cur_cert = str(ws.Range(cert_cell).Value or "").strip().replace('"', '""')
        cert_idx = patron_lookup(id_cell, "F")
        ws.Range(cert_cell).Formula = (
            f'=IF(IFERROR({cert_idx},"")="","{cur_cert}",{cert_idx})'
        )
        cur_vig = ws.Range(vig_cell).Value
        if cur_vig is not None and hasattr(cur_vig, "year"):
            fallback = f"DATE({cur_vig.year},{cur_vig.month},{cur_vig.day})"
        else:
            fallback = '""'
            try:
                if ws.Range(vig_cell).HasFormula:
                    fallback_formula = ws.Range(vig_cell).Formula
                    vig_idx = patron_lookup(id_cell, "H")
                    ws.Range(vig_cell).Formula = (
                        f'=IF(IFERROR({vig_idx},"")="",{fallback_formula[1:]},{vig_idx})'
                    )
                    wired += 1
                    continue
            except Exception:
                pass
        vig_idx = patron_lookup(id_cell, "H")
        ws.Range(vig_cell).Formula = (
            f'=IF(IFERROR({vig_idx},"")="",{fallback},{vig_idx})'
        )
        wired += 1
    return wired


VBA_CODE = r'''
Option Explicit

Private Const AG_PASSWORD As String = "AG-Calidad-2026"
Private Const MM_PER_IN As Double = 25.4

Private Function CertificadoDim() As String
    With ThisWorkbook.Worksheets("Calculos")
        CertificadoDim = Trim(CStr(.Range("D4").Value)) & "-" & _
                         Format(.Range("E4").Value, "0000") & "-" & _
                         Format(.Range("F4").Value, "00")
    End With
End Function

Private Sub LimpiarLecturas(ws As Worksheet)
    Dim r As Long
    For r = 18 To 28
        ws.Range("C" & r & ":E" & r).ClearContents
    Next r
End Sub

Private Sub ConvertirEmp(ws As Worksheet, factor As Double)
    Dim r As Long
    Dim v As Variant
    For r = 18 To 28
        v = ws.Range("H" & r).Value
        If IsNumeric(v) Then
            On Error Resume Next
            If Not ws.Range("H" & r).HasFormula Then
                ws.Range("H" & r).Value = CDbl(v) * factor
            End If
            On Error GoTo 0
        End If
    Next r
End Sub

Sub GuardarCertificadoExcel()
    Dim ws As Worksheet
    Dim ruta As Variant
    Dim nombreArchivo As String
    Dim instrumento As String
    Dim idEquipo As String
    Dim avisoPatron As String

    On Error GoTo ErrorHandler
    Set ws = ThisWorkbook.Worksheets("Calculos")

    If Trim(CStr(ws.Range("D4").Value)) = "" Or _
       Trim(CStr(ws.Range("E4").Value)) = "" Or _
       Trim(CStr(ws.Range("F4").Value)) = "" Then
        MsgBox "El número de certificado está incompleto (D4-E4-F4).", _
               vbCritical, "Validación"
        Exit Sub
    End If

    instrumento = Trim(CStr(ws.Range("B9").Value))
    idEquipo = Trim(CStr(ws.Range("F9").Value))
    If instrumento = "" Or instrumento = "No encontrado" Or idEquipo = "" Then
        MsgBox "Falta el instrumento o número de control. Revisa el certificado o pulsa Actualizar.", _
               vbCritical, "Validación"
        Exit Sub
    End If

    On Error Resume Next
    avisoPatron = CStr(ThisWorkbook.Worksheets("PORTADA").Range("E35").Value)
    On Error GoTo ErrorHandler
    If InStr(1, UCase(avisoPatron), "VENCIDO") > 0 Then
        If MsgBox("Un patrón aparece VENCIDO." & vbCrLf & "¿Deseas guardar de todos modos?", _
                  vbExclamation + vbYesNo, "Calidad") = vbNo Then Exit Sub
    End If

    nombreArchivo = CertificadoDim() & " - " & instrumento & " - " & idEquipo
    nombreArchivo = Replace(nombreArchivo, "/", "-")
    nombreArchivo = Replace(nombreArchivo, "\", "-")
    nombreArchivo = Replace(nombreArchivo, ":", "")
    nombreArchivo = Replace(nombreArchivo, "*", "")
    nombreArchivo = Replace(nombreArchivo, "?", "")
    nombreArchivo = Replace(nombreArchivo, """", "")
    nombreArchivo = Replace(nombreArchivo, "<", "")
    nombreArchivo = Replace(nombreArchivo, ">", "")
    nombreArchivo = Replace(nombreArchivo, "|", "")

    ruta = Application.GetSaveAsFilename( _
        InitialFileName:=nombreArchivo, _
        FileFilter:="Libro de Excel con macros (*.xlsm), *.xlsm", _
        Title:="Guardar certificado de regla / flexómetro")

    If ruta = False Then Exit Sub

    Application.DisplayAlerts = False
    ThisWorkbook.SaveCopyAs CStr(ruta)
    Application.DisplayAlerts = True
    MsgBox "Certificado guardado:" & vbCrLf & CStr(ruta), vbInformation, "Listo"
    Exit Sub

ErrorHandler:
    Application.DisplayAlerts = True
    MsgBox "No se pudo guardar: " & Err.Description, vbCritical, "Error"
End Sub

Sub CambiarUnidadReglaFlex()
    Dim ws As Worksheet
    Dim unidad As String
    Dim alcance As Variant
    Dim divMin As Variant

    Set ws = ThisWorkbook.Worksheets("Calculos")
    On Error Resume Next
    ws.Unprotect Password:=AG_PASSWORD
    On Error GoTo 0

    unidad = LCase(Trim(CStr(ws.Range("J11").Value)))
    alcance = ws.Range("F10").Value
    divMin = ws.Range("J9").Value

    If unidad = "mm" Then
        ws.Range("J11").Value = "in"
        If IsNumeric(alcance) Then ws.Range("F10").Value = CDbl(alcance) / MM_PER_IN
        If IsNumeric(divMin) Then ws.Range("J9").Value = CDbl(divMin) / MM_PER_IN
        ConvertirEmp ws, 1# / MM_PER_IN
        LimpiarLecturas ws
        MsgBox "Modo pulgadas (in)." & vbCrLf & _
               "Alcance/división/EMP convertidos ÷25.4." & vbCrLf & _
               "Intervalos % se recalculan; captura de nuevo las lecturas.", _
               vbInformation, "Unidad"
    Else
        ws.Range("J11").Value = "mm"
        If IsNumeric(alcance) Then ws.Range("F10").Value = CDbl(alcance) * MM_PER_IN
        If IsNumeric(divMin) Then ws.Range("J9").Value = CDbl(divMin) * MM_PER_IN
        ConvertirEmp ws, MM_PER_IN
        LimpiarLecturas ws
        MsgBox "Modo milímetros (mm)." & vbCrLf & _
               "Alcance/división/EMP convertidos ×25.4." & vbCrLf & _
               "Intervalos % se recalculan; captura de nuevo las lecturas.", _
               vbInformation, "Unidad"
    End If

    On Error Resume Next
    ws.Protect Password:=AG_PASSWORD, DrawingObjects:=False, Contents:=True, Scenarios:=True
    Application.Calculate
    On Error GoTo 0
End Sub

Sub RecalcularCertificado()
    On Error GoTo ErrorHandler
    Application.ScreenUpdating = False
    ThisWorkbook.RefreshAll
    Application.CalculateUntilAsyncQueriesDone
    ThisWorkbook.Worksheets("Calculos").Calculate
    ThisWorkbook.Worksheets("PORTADA").Calculate
    ThisWorkbook.Worksheets("RESULTADOS").Calculate
    Application.ScreenUpdating = True
    MsgBox "Datos y cálculos actualizados para " & CertificadoDim(), _
           vbInformation, "Actualizado"
    Exit Sub

ErrorHandler:
    Application.ScreenUpdating = True
    MsgBox "No se pudo actualizar: " & Err.Description, vbExclamation, "Actualizar"
End Sub

Sub IrAPortada()
    ThisWorkbook.Worksheets("PORTADA").Activate
    ThisWorkbook.Worksheets("PORTADA").Range("A1").Select
End Sub
'''


def rgb(red: int, green: int, blue: int) -> int:
    return red + green * 256 + blue * 65536


def set_module(vb_project, name: str, code_text: str) -> None:
    component = None
    for index in range(1, vb_project.VBComponents.Count + 1):
        candidate = vb_project.VBComponents.Item(index)
        if str(candidate.Name) == name:
            component = candidate
            break
    if component is None:
        component = vb_project.VBComponents.Add(1)
        component.Name = name
    code = component.CodeModule
    if code.CountOfLines:
        code.DeleteLines(1, code.CountOfLines)
    code.AddFromString(code_text)


def add_button(ws, left: float, top: float, width: float, caption: str, macro: str, color: int) -> None:
    shape = ws.Shapes.AddShape(MSO_ROUNDED, left, top, width, 25)
    shape.Name = f"btn_{macro}"
    shape.OnAction = macro
    shape.Fill.ForeColor.RGB = color
    shape.Line.Visible = 0
    shape.Placement = 3
    shape.Locked = True
    shape.TextFrame.Characters().Text = caption
    font = shape.TextFrame.Characters().Font
    font.Color = 0xFFFFFF
    font.Bold = True
    font.Size = 9
    font.Name = "Calibri"
    shape.TextFrame.HorizontalAlignment = -4108
    shape.TextFrame.VerticalAlignment = -4108


def delete_query_if_exists(wb, name: str) -> None:
    try:
        wb.Queries.Item(name).Delete()
    except Exception:
        pass


def delete_connection_if_exists(wb, name: str) -> None:
    try:
        wb.Connections.Item(name).Delete()
    except Exception:
        pass


def ensure_sheet(wb, name: str):
    try:
        return wb.Worksheets(name)
    except Exception:
        ws = wb.Worksheets.Add(After=wb.Worksheets(wb.Worksheets.Count))
        ws.Name = name[:31]
        return ws


def clear_sheet(ws) -> None:
    while ws.ListObjects.Count > 0:
        ws.ListObjects(1).Delete()
    ws.Cells.Clear()


def load_query_to_sheet(wb, query_name: str, ws, table_name: str) -> None:
    clear_sheet(ws)
    conn_name = f"Query - {query_name}"
    delete_connection_if_exists(wb, conn_name)
    source = (
        "OLEDB;Provider=Microsoft.Mashup.OleDb.1;"
        f'Data Source=$Workbook$;Location="{query_name}";Extended Properties=""'
    )
    table = ws.ListObjects.Add(0, source, False, 1, ws.Range("A1"))
    query = table.QueryTable
    query.CommandType = 2
    query.CommandText = [f"SELECT * FROM [{query_name}]"]
    query.BackgroundQuery = False
    query.PreserveFormatting = True
    query.RefreshOnFileOpen = True
    query.RefreshPeriod = 0
    query.Refresh(False)
    table.Name = table_name
    table.DisplayName = table_name
    connection = None
    try:
        connection = wb.Connections.Item(conn_name)
    except Exception:
        for index in range(1, wb.Connections.Count + 1):
            candidate = wb.Connections.Item(index)
            if query_name.lower() in str(candidate.Name).lower():
                connection = candidate
                break
    if connection is not None:
        connection.RefreshWithRefreshAll = True
        try:
            connection.OLEDBConnection.BackgroundQuery = False
            connection.OLEDBConnection.RefreshOnFileOpen = True
            connection.OLEDBConnection.RefreshPeriod = 0
        except Exception:
            pass


def try_unprotect(ws) -> None:
    for pwd in (PASSWORD, "AG", "calidad", "1234", ""):
        try:
            if pwd:
                ws.Unprotect(Password=pwd)
            else:
                ws.Unprotect()
            return
        except Exception:
            continue


def try_unmerge(ws, addr: str) -> None:
    try:
        rng = ws.Range(addr)
        if rng.MergeCells:
            rng.MergeArea.UnMerge()
    except Exception:
        pass


def unlock_cell(ws, addr: str) -> None:
    try:
        rng = ws.Range(addr)
        if rng.MergeCells:
            rng = rng.MergeArea.Cells(1, 1)
        rng.Locked = False
        rng.Interior.Color = rgb(255, 242, 204)
    except Exception:
        try:
            ws.Range(addr).Interior.Color = rgb(255, 242, 204)
        except Exception:
            pass


def strip_sheet_protection(src: Path, dst: Path) -> int:
    removed = 0
    pattern = re.compile(
        rb"<sheetProtection\b[^>]*/>|<sheetProtection\b[\s\S]*?</sheetProtection>",
        re.I,
    )
    with zipfile.ZipFile(src, "r") as zin:
        buf = BytesIO()
        with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename.startswith("xl/worksheets/") and info.filename.endswith(".xml"):
                    new_data, n = pattern.subn(b"", data)
                    if n:
                        removed += n
                        data = new_data
                zout.writestr(info, data)
        dst.write_bytes(buf.getvalue())
    return removed


def formula_m(row: int) -> str:
    """uErrPat dual mm/in; corrige precedencia /25.4 e IFERROR."""
    return (
        f'=IF($J$11="mm",'
        f'IFERROR(VLOOKUP(B{row},Patrones!$E$6:$H$16,4,TRUE),0)'
        f'+IFERROR(VLOOKUP(G{row},Patrones!$E$20:$H$29,4,TRUE),0),'
        f'(IFERROR(VLOOKUP(B{row}*25.4,Patrones!$E$6:$H$16,4,TRUE),0)'
        f'+IFERROR(VLOOKUP(G{row}*25.4,Patrones!$E$20:$H$29,4,TRUE),0))/25.4)'
    )


def formula_o(row: int) -> str:
    """uPat dual; reticula U fija $G$20 (evita G30 vacío en fila 28)."""
    return (
        f'=IF($J$11="mm",'
        f'IFERROR(VLOOKUP(B{row},Patrones!$E$6:$H$16,3,TRUE),0)+IFERROR(Patrones!$G$20,0),'
        f'(IFERROR(VLOOKUP(B{row}*25.4,Patrones!$E$6:$H$16,3,TRUE),0)'
        f'+IFERROR(Patrones!$G$20,0))/25.4)'
    )


def apply_dual_unit_fixes(calc) -> None:
    calc.Range("J11").Value = "mm"
    calc.Range("A18").Value = 0
    calc.Range("B18").Value = 0
    calc.Range("B19").Formula = "=$F$10/10"

    for row in range(19, 29):
        # Intervalo A sigue al valor deseado / alcance (mm e in)
        calc.Range(f"A{row}").Formula = f"=B{row}"

    for row in range(18, 29):
        calc.Range(f"M{row}").Formula = formula_m(row)
        calc.Range(f"O{row}").Formula = formula_o(row)
        calc.Range(f"K{row}").Formula = f"=IFERROR(STDEV(C{row}:E{row}),0)"
        calc.Range(f"J{row}").Formula = (
            f'=IFERROR(IF(ABS(G{row})+T{row}>H{row},"RECHAZADO","ACEPTADO"),'
            f'IF(ABS(G{row})>H{row},"RECHAZADO","ACEPTADO"))'
        )
        calc.Range(f"I{row}").Formula = (
            f'=IFERROR(IF(ABS(G{row})>H{row},"RECHAZADO","ACEPTADO"),"")'
        )
        calc.Range(f"F{row}").Formula = (
            f'=IF(COUNT(C{row}:E{row})<1,"",AVERAGE(C{row}:E{row}))'
        )
        calc.Range(f"G{row}").Formula = f'=IF(F{row}="","",F{row}-A{row})'


def delete_sheet_if_exists(wb, name: str) -> None:
    try:
        ws = wb.Worksheets(name)
    except Exception:
        return
    try_unprotect(ws)
    excel_app = wb.Application
    excel_app.DisplayAlerts = False
    ws.Delete()
    excel_app.DisplayAlerts = False


def copy_inch_patrones_from_in(patrones_mm, source_in: Path) -> None:
    """Copia columnas I:L (resultados in) desde el xlsx in si existen."""
    if not source_in.exists():
        return
    try:
        import openpyxl
        wb_in = openpyxl.load_workbook(source_in, data_only=False)
        pat_in = wb_in["Patrones"]
        for row in range(4, 30):
            for col in range(9, 13):  # I-L
                val = pat_in.cell(row, col).value
                if val is None:
                    continue
                addr = f"{chr(64 + col)}{row}" if col <= 26 else None
                if addr is None:
                    continue
                if isinstance(val, str) and val.startswith("="):
                    patrones_mm.Range(addr).Formula = val
                else:
                    patrones_mm.Range(addr).Value = val
        wb_in.close()
        print("  Columnas in (I:L) de Patrones copiadas desde xlsx in")
    except Exception as exc:
        print(f"  (aviso) No se copiaron columnas in de Patrones: {exc}")


def main() -> int:
    if not SOURCE_MM.exists():
        print(f"No existe: {SOURCE_MM}")
        return 1

    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    work = FOLDER / f"_tmp_regla_flex_{stamp}.xlsx"
    backup_mm = FOLDER / f"Formato regla-flex mm_backup_setup_{stamp}.xlsx"
    shutil.copy2(SOURCE_MM, backup_mm)
    print(f"Respaldo mm: {backup_mm.name}")
    if SOURCE_IN.exists():
        backup_in = FOLDER / f"Formato regla-flex in_backup_setup_{stamp}.xlsx"
        shutil.copy2(SOURCE_IN, backup_in)
        print(f"Respaldo in: {backup_in.name}")

    n_prot = strip_sheet_protection(SOURCE_MM, work)
    print(f"Trabajo: {work.name} (protección quitada en {n_prot} hojas)")

    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    excel.AskToUpdateLinks = False
    excel.EnableEvents = False
    wb = None

    try:
        print(f"Abriendo: {work.name}")
        wb = excel.Workbooks.Open(str(work.resolve()), UpdateLinks=0, ReadOnly=False)
        if wb.ReadOnly:
            raise RuntimeError("Solo lectura")
        try:
            wb.Queries.FastCombine = True
        except Exception:
            pass

        calc = wb.Worksheets("Calculos")
        portada = wb.Worksheets("PORTADA")
        resultados = wb.Worksheets("RESULTADOS")
        patrones = wb.Worksheets("Patrones")
        for ws in (calc, portada, resultados, patrones):
            try_unprotect(ws)
        try:
            try_unprotect(wb.Worksheets("CMC"))
        except Exception:
            pass

        delete_sheet_if_exists(wb, "AG-ADM-F10-00")
        print("Eliminada hoja AG-ADM-F10-00")

        copy_inch_patrones_from_in(patrones, SOURCE_IN)

        query_defs = (
            ("AG_API_Historial_ReglaFlex", M_HISTORIAL, HIST_SHEET, "AG_Historial"),
            ("AG_API_Clientes_ReglaFlex", M_CLIENTES, CLIENTES_SHEET, "AG_Clientes"),
            ("AG_API_Patrones_ReglaFlex", M_PATRONES, PATRONES_SHEET, "AG_Patrones"),
        )
        print("Creando Power Query…")
        for query_name, _, _, _ in query_defs:
            delete_query_if_exists(wb, query_name)
            delete_connection_if_exists(wb, f"Query - {query_name}")
        for query_name, formula, sheet_name, table_name in query_defs:
            print(f"  {query_name} -> {sheet_name}")
            wb.Queries.Add(query_name, formula)
            load_query_to_sheet(wb, query_name, ensure_sheet(wb, sheet_name), table_name)

        print("Cableando certificado e historial…")
        for ref in ("A4", "C4", "D4", "E4", "F4"):
            try_unmerge(calc, ref)
        calc.Range("A4").Value = "****No. DE CERTIFICADO:"
        try:
            calc.Range("A4:C4").Merge()
        except Exception:
            pass
        calc.Range("D4").Value = "AGD"
        calc.Range("E4").Value = 1278
        calc.Range("F4").Value = 23
        for ref in ("D4", "E4", "F4", "J11", "F10", "J9", "I2", "I5"):
            unlock_cell(calc, ref)
        for row in range(18, 29):
            unlock_cell(calc, f"C{row}")
            unlock_cell(calc, f"D{row}")
            unlock_cell(calc, f"E{row}")
            unlock_cell(calc, f"H{row}")

        for addr, formula in CALC_FORMULAS.items():
            calc.Range(addr).Formula = formula

        calc.Range("AH4").Font.Size = 8
        calc.Range("AH4").Font.Color = rgb(150, 150, 150)

        portada.Range("J10").Formula = F_CERT
        resultados.Range("J10").Formula = F_CERT

        print("Corrigiendo fórmulas duales mm/in e intervalos…")
        apply_dual_unit_fixes(calc)

        n_wired = wire_patron_blocks(wb)
        print(f"Patrones enlazados: {n_wired}")

        print("Macros y botones…")
        try:
            set_module(wb.VBProject, "ModuloAG_ReglaFlexUI", VBA_CODE)
        except Exception as exc:
            raise RuntimeError("Activa confianza al modelo de objetos VBA.") from exc

        old = [
            str(calc.Shapes(i).Name)
            for i in range(1, calc.Shapes.Count + 1)
            if str(calc.Shapes(i).Name).startswith("btn_")
        ]
        for name in old:
            calc.Shapes(name).Delete()

        buttons = (
            ("Guardar", "GuardarCertificadoExcel", 92, rgb(37, 99, 235)),
            ("mm / in", "CambiarUnidadReglaFlex", 88, rgb(124, 58, 237)),
            ("Actualizar", "RecalcularCertificado", 88, rgb(217, 119, 6)),
            ("Ir a Portada", "IrAPortada", 88, rgb(71, 85, 105)),
        )
        left = float(calc.Range("V1").Left) + 6
        top = 4.0
        for caption, macro, width, color in buttons:
            add_button(calc, left, top, width, caption, macro, color)
            left += width + 5

        for sheet_name in (HIST_SHEET, CLIENTES_SHEET, PATRONES_SHEET):
            wb.Worksheets(sheet_name).Visible = XL_VERY_HIDDEN

        excel.CalculateUntilAsyncQueriesDone()
        excel.Calculate()
        time.sleep(1)

        # Prueba rápida mm
        calc.Range("J11").Value = "mm"
        calc.Range("F10").Value = 5000
        excel.Calculate()
        a28 = calc.Range("A28").Value
        b19 = calc.Range("B19").Value
        print(f"  Prueba mm/5000: B19={b19} A28={a28} M18={calc.Range('M18').Value}")

        # Prueba in
        calc.Range("J11").Value = "in"
        calc.Range("F10").Value = 144
        excel.Calculate()
        print(
            f"  Prueba in/144: B19={calc.Range('B19').Value} "
            f"A28={calc.Range('A28').Value} M19={calc.Range('M19').Value}"
        )
        calc.Range("J11").Value = "mm"
        calc.Range("F10").Value = 5000
        calc.Range("J9").Value = 0.5

        calc.Protect(Password=PASSWORD, DrawingObjects=False, Contents=True, Scenarios=True)
        portada.Protect(Password=PASSWORD, DrawingObjects=False, Contents=True, Scenarios=True)
        patrones.Protect(Password=PASSWORD, DrawingObjects=False, Contents=True, Scenarios=True)

        if TARGET.exists():
            TARGET.unlink()
        wb.SaveAs(str(TARGET.resolve()), FileFormat=XL_XLSM)
        print(f"Guardado: {TARGET.name}")

        if not bool(wb.HasVBProject):
            raise RuntimeError("Sin VBA")
        found = {
            str(calc.Shapes(i).Name)
            for i in range(1, calc.Shapes.Count + 1)
            if str(calc.Shapes(i).Name).startswith("btn_")
        }
        if len(found) != 4:
            raise RuntimeError(f"Botones: {found}")

        print(f"  Cert Portada: {portada.Range('J10').Value}")
        print(f"  Unidad: {calc.Range('J11').Value} Alcance: {calc.Range('F10').Value}")

        wb.Close(SaveChanges=True)
        wb = None

        for src in (SOURCE_MM, SOURCE_IN):
            if src.exists():
                src.unlink()
                print(f"Retirado: {src.name}")

        print("LISTO: Formato Regla Flex.xlsm")
        return 0
    except Exception as exc:
        print(f"ERROR: {exc}")
        import traceback
        traceback.print_exc()
        return 1
    finally:
        if wb is not None:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        try:
            excel.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()
        if work.exists():
            try:
                work.unlink()
            except Exception:
                pass


if __name__ == "__main__":
    raise SystemExit(main())
